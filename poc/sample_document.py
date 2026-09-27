"""Genere un document .docx synthetique (fiche technique) pour le PoC.

Contenu invente pour ce test, sans rapport avec un vrai document du creator.
Structure volontairement proche de l'exemple donne dans la conception :
Titre, paragraphe de resume, section "Caracteristiques" avec un tableau,
section "Normes".
"""
from __future__ import annotations

from pathlib import Path

from docx import Document

OUT = Path(__file__).resolve().parent / "out" / "fiche_technique_test.docx"


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.title = "Fiche technique - Joint torique EPDM 45x3"
    doc.core_properties.author = "Test PoC IAFActory"
    doc.core_properties.subject = "Fiche technique materiau"

    doc.add_heading("Fiche technique - Joint torique EPDM 45x3", level=0)

    doc.add_heading("Resume", level=1)
    doc.add_paragraph(
        "Ce joint torique en EPDM est destine aux applications hydrauliques "
        "basse pression. Il resiste aux temperatures de -40 a 120 degres."
    )

    doc.add_heading("Caracteristiques", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Propriete"
    hdr[1].text = "Valeur"
    rows = [
        ("Materiau", "EPDM"),
        ("Diametre interieur", "45 mm"),
        ("Epaisseur", "3 mm"),
        ("Resistance a la traction", "9 MPa"),
        ("Temperature d'utilisation", "-40 a 120 degres C"),
        ("Fournisseur", "Acme Joints SA"),
    ]
    for prop, val in rows:
        cells = table.add_row().cells
        cells[0].text = prop
        cells[1].text = val

    doc.add_heading("Normes", level=1)
    doc.add_paragraph(
        "Ce joint est conforme a la norme ISO 3601-1 classe A. Il est "
        "egalement compatible avec les exigences de la norme DIN 3771."
    )

    doc.save(OUT)
    print(f"Document genere : {OUT}")


if __name__ == "__main__":
    build()
