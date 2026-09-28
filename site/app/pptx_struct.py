"""Analyse structurelle minimale d'un .pptx (US3.1, US3.8). Chaque diapositive
devient une Section (titre du placeholder titre si present, sinon
"Diapositive N") ; chaque paragraphe non vide d'une zone de texte (hors la
zone titre, deja utilisee comme label de section) devient un Paragraph ;
chaque tableau devient un Table. Pas de rendu des images ni des schemas
(legendes seulement si elles sont du texte, US3.8 - description par modele a
vision hors scope ici)."""
from __future__ import annotations

import itertools

from pptx import Presentation

from .struct_element import StructElement


def parse(path: str) -> tuple[dict, StructElement]:
    prs = Presentation(path)
    cp = prs.core_properties
    metadata = {"title": cp.title, "author": cp.author, "subject": cp.subject}

    root = StructElement(kind="Document", label="racine", level=0, position=0)
    position = itertools.count(1)

    for slide_num, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title_text = (title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else "")
        section = StructElement(
            kind="Section", label=title_text or f"Diapositive {slide_num}", level=1, position=next(position),
        )
        root.children.append(section)

        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue  # deja utilise comme label de section
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if not text:
                        continue
                    elem = StructElement(
                        kind="Paragraph", label=text[:60], level=2, position=next(position), text=text,
                    )
                    section.children.append(elem)
            if shape.has_table:
                rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in shape.table.rows]
                table_text = "\n".join(rows)
                if not table_text.strip():
                    continue
                elem = StructElement(
                    kind="Table", label=f"Tableau ({len(shape.table.rows)}x{len(shape.table.columns)})",
                    level=2, position=next(position), text=table_text,
                )
                section.children.append(elem)

    return metadata, root
