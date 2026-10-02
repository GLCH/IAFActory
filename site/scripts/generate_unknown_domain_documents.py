"""Regenere les documents Word des domaines INCONNUS du jeu d'essai
(`corpus/documents/inconnus/`) : astronomie, cuisine, droit des contrats.

Contrat (comme generate_llm_document.py) : Gemini REDIGE le texte via la
passerelle LLM (`app.graph.chat`), donc le contenu n'est pas identique d'un
lancement a l'autre. Les fichiers deja presents dans corpus/ sont la version
de reference du scenario enregistre le 2026-10-02 ; ce script sert a en
produire de NOUVEAUX (autre graine de contenu, autres sujets).

Usage (depuis site/) :
  python scripts/generate_unknown_domain_documents.py --out ../corpus/documents/inconnus/nouveaux
  python scripts/generate_unknown_domain_documents.py --out /tmp/x --only astro-1,droit-1
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document  # noqa: E402

from app.config import settings  # noqa: E402
from app.graph import chat  # noqa: E402

SYSTEM = (
    "Tu rediges en francais un document structure (3 sections courtes: Introduction, Developpement, Conclusion), "
    "factuel et specialise, 250 a 320 mots au total. Reponds par le texte seul, avec les titres de section "
    "sur leur propre ligne precedes de '## '."
)

TOPICS = {
    "cuisine-1": "une recette de gratin dauphinois : pommes de terre, creme, ail, cuisson au four, ingredients et temps de preparation",
    "cuisine-2": "une recette de tarte aux pommes : pate brisee, pommes, sucre, beurre, cuisson, ingredients et astuces de patissier",
    "cuisine-3": "une recette de ratatouille provencale : aubergines, courgettes, poivrons, tomates, huile d'olive, cuisson mijotee",
    "cuisine-4": "une recette de soupe a l'oignon gratinee : oignons caramelises, bouillon, pain, fromage, gratinage",
    "astro-1": "les trous noirs en astrophysique : horizon des evenements, singularite, masse stellaire, ondes gravitationnelles, rayonnement de Hawking",
    "astro-2": "les planetes du systeme solaire : orbites, planetes telluriques et geantes gazeuses, satellites naturels, ceinture d'asteroides",
    "astro-3": "les galaxies et la cosmologie : Voie lactee, expansion de l'univers, matiere noire, Big Bang, fond diffus cosmologique",
    "astro-4": "les telescopes et l'observation du ciel : optique, spectroscopie, telescopes spatiaux, radioastronomie, magnitude des etoiles",
    "astro-5": "la vie des etoiles : nebuleuse, fusion nucleaire, sequence principale, geante rouge, supernova, naine blanche, etoile a neutrons",
    "droit-1": "le droit des contrats : formation du contrat, consentement, clauses, responsabilite contractuelle, resiliation, tribunal competent",
}


def write_document(name: str, text: str, out_dir: Path) -> Path:
    doc = Document()
    doc.add_heading(name.replace("-", " ").title(), level=0)
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("##"):
            doc.add_heading(line.lstrip("# ").strip(), level=1)
        else:
            doc.add_paragraph(line)
    path = out_dir / f"{name}.docx"
    doc.save(str(path))
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--out", required=True, help="dossier de sortie (cree si besoin)")
    parser.add_argument("--only", help="noms separes par des virgules (defaut : tous)")
    args = parser.parse_args()
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    names = [n.strip() for n in args.only.split(",")] if args.only else list(TOPICS)
    unknown = [n for n in names if n not in TOPICS]
    if unknown:
        raise SystemExit(f"sujets inconnus : {unknown} (disponibles : {list(TOPICS)})")
    for name in names:
        text = chat(f"Redige le document sur : {TOPICS[name]}.", model=settings.answer_model, system=SYSTEM, max_tokens=1500)
        path = write_document(name, text, out_dir)
        print(path.name, len(text.split()), "mots")


if __name__ == "__main__":
    main()
